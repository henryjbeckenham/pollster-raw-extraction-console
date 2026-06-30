from __future__ import annotations

import csv
import json
import zipfile
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Sequence

from openpyxl import Workbook

from extractors.pdf_extractor import ExtractedTable, ExtractionResult
from utils.file_naming import unique_folder_names


def build_export_zip(results: Sequence[ExtractionResult]) -> bytes:
    folder_names = unique_folder_names([result.source_filename for result in results])

    with TemporaryDirectory(prefix="pollster_raw_extraction_") as temp_dir:
        temp_path = Path(temp_dir)

        for result, folder_name in zip(results, folder_names, strict=True):
            output_dir = temp_path / folder_name
            write_report_outputs(result, output_dir)

        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for file_path in sorted(temp_path.rglob("*")):
                if file_path.is_file():
                    archive.write(file_path, file_path.relative_to(temp_path).as_posix())

    return zip_buffer.getvalue()


def write_report_outputs(result: ExtractionResult, output_dir: Path) -> list[str]:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_dir = output_dir / "tables_csv"
    csv_dir.mkdir(exist_ok=True)
    image_dir = output_dir / "page_images"
    image_dir.mkdir(exist_ok=True)

    output_files = _expected_output_files(result)

    _write_report_text(result, output_dir / "report_text.md")
    _write_report_raw_json(result, output_dir / "report_raw.json")
    _write_table_csvs(result, csv_dir)
    _write_page_images(result, image_dir)
    _write_report_tables_workbook(result, output_dir / "report_tables.xlsx")
    _write_manifest(result, output_dir / "manifest.json", output_files)

    return output_files


def _expected_output_files(result: ExtractionResult) -> list[str]:
    output_files = [
        "manifest.json",
        "report_text.md",
        "report_tables.xlsx",
        "report_raw.json",
    ]

    output_files.extend(
        f"tables_csv/{table.csv_filename}" for table in _iter_tables(result)
    )
    output_files.extend(
        page.page_image_filename for page in result.pages if page.image_png
    )
    return output_files


def _write_report_text(result: ExtractionResult, output_path: Path) -> None:
    lines = [
        f"# {result.source_filename}",
        "",
        f"Source filename: {result.source_filename}",
        f"Page count: {result.page_count}",
        "",
    ]

    for page in result.pages:
        lines.extend(
            [
                f"## Page {page.page_number}",
                "",
                page.text.rstrip() or "_No text extracted._",
                "",
            ]
        )

    output_path.write_text("\n".join(lines), encoding="utf-8")


def _write_report_raw_json(result: ExtractionResult, output_path: Path) -> None:
    output_path.write_text(
        json.dumps(result.to_raw_dict(), indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _write_table_csvs(result: ExtractionResult, csv_dir: Path) -> None:
    for table in _iter_tables(result):
        with (csv_dir / table.csv_filename).open("w", newline="", encoding="utf-8") as csv_file:
            writer = csv.writer(csv_file)
            writer.writerows(table.rows)


def _write_page_images(result: ExtractionResult, image_dir: Path) -> None:
    for page in result.pages:
        if page.image_png:
            (image_dir / Path(page.page_image_filename).name).write_bytes(page.image_png)


def _write_report_tables_workbook(result: ExtractionResult, output_path: Path) -> None:
    workbook = Workbook()
    table_index_sheet = workbook.active
    table_index_sheet.title = "Table_Index"

    table_index_headers = [
        "table_id",
        "page_number",
        "table_number_on_page",
        "row_count",
        "column_count",
        "csv_filename",
        "notes",
    ]
    table_index_sheet.append(table_index_headers)

    for table in _iter_tables(result):
        table_index_sheet.append([_table_index_row(table)[header] for header in table_index_headers])

        table_sheet = workbook.create_sheet(title=table.table_id)
        for row in table.rows:
            table_sheet.append(row)

    workbook.save(output_path)


def _write_manifest(
    result: ExtractionResult,
    output_path: Path,
    output_files: list[str],
) -> None:
    manifest = {
        "original_filename": result.source_filename,
        "extraction_timestamp": result.extraction_timestamp,
        "output_files_created": output_files,
        "page_image_files": [
            page.page_image_filename for page in result.pages if page.image_png
        ],
        "total_page_count": result.page_count,
        "total_extracted_tables": result.total_tables,
        "warnings": result.warnings,
        "errors": result.errors,
    }

    output_path.write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def _table_index_row(table: ExtractedTable) -> dict[str, object]:
    return {
        "table_id": table.table_id,
        "page_number": table.page_number,
        "table_number_on_page": table.table_number_on_page,
        "row_count": table.row_count,
        "column_count": table.column_count,
        "csv_filename": table.csv_filename,
        "notes": table.notes,
    }


def _iter_tables(result: ExtractionResult) -> list[ExtractedTable]:
    return [table for page in result.pages for table in page.tables]
